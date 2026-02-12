import os
from docx import Document
from fpdf import FPDF

RAW_DIR = "anap_rag_pipeline/data/raw/mock_cell"
os.makedirs(RAW_DIR, exist_ok=True)

def create_docx():
    doc = Document()
    doc.add_heading("Contract de Servicii - Model Standard", 0)

    doc.add_heading("Sectiunea I - Partile Contractante", 1)
    doc.add_paragraph("Prezentul contract se incheie intre Autoritatea Contractanta si Prestator.")

    doc.add_heading("Sectiunea II - Obiectul Contractului", 1)
    doc.add_paragraph("Obiectul prezentului contract il constituie prestarea serviciilor de curatenie.")

    doc.add_heading("Sectiunea III - Pretul si Modalitatile de Plata", 1)
    doc.add_paragraph("Pretul contractului este ferm in lei.")

    # Add a table
    table = doc.add_table(rows=1, cols=3)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Nr. Crt.'
    hdr_cells[1].text = 'Descriere Serviciu'
    hdr_cells[2].text = 'Pret Unitar'

    row_cells = table.add_row().cells
    row_cells[0].text = '1'
    row_cells[1].text = 'Curatenie birouri'
    row_cells[2].text = '100 RON/mp'

    doc.save(os.path.join(RAW_DIR, "mock_contract.docx"))
    print(f"Created {os.path.join(RAW_DIR, 'mock_contract.docx')}")

def create_pdf():
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)

    # Header repetition simulation (ASCII only to avoid font issues)
    header = "MINISTERUL INVESTITIILOR SI PROIECTELOR EUROPENE\nDirectia Generala Achizitii Publice\n"

    pdf.multi_cell(0, 10, header, align='C')
    pdf.ln(10)

    pdf.set_font("Arial", 'B', 14)
    pdf.cell(0, 10, "FISA DE DATE A ACHIZITIEI", 0, 1, 'C')
    pdf.ln(5)

    pdf.set_font("Arial", size=11)
    content = """
    SECTIUNEA I: AUTORITATEA CONTRACTANTA
    I.1) Denumire si adrese
    Denumire oficiala: Agentia Nationala pentru Achizitii Publice

    SECTIUNEA II: OBIECTUL
    II.1.1) Titlu: Servicii de consultanta in management

    Lorem ipsum dolor sit amet, consectetur adipiscing elit. Sed do eiusmod tempor incididunt ut labore et dolore magna aliqua.
    Ut enim ad minim veniam, quis nostrud exercitation ullamco laboris nisi ut aliquip ex ea commodo consequat.
    """
    pdf.multi_cell(0, 5, content)

    # Page 2 with same header
    pdf.add_page()
    # Reset font for header
    pdf.set_font("Arial", size=12)
    pdf.multi_cell(0, 10, header, align='C')
    pdf.ln(10)
    pdf.set_font("Arial", size=11)
    pdf.multi_cell(0, 5, "Continuare continut Fisa de Date...\nMai multe detalii tehnice aici.")

    pdf.output(os.path.join(RAW_DIR, "mock_fisa_date.pdf"))
    print(f"Created {os.path.join(RAW_DIR, 'mock_fisa_date.pdf')}")

if __name__ == "__main__":
    create_docx()
    try:
        create_pdf()
    except Exception as e:
        print(f"Skipping PDF creation: {e}")
